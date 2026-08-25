#!/usr/bin/env python3
"""
Enrichissement complet des établissements depuis le XLSX.

Cible :
- fiches_etablissement (colonnes ajoutées par migration_etablissement_xlsx.sql)
- Tables satellites :
  * etablissement_reseaux_sociaux
  * etablissement_admissions
  * etablissement_responsables
  * etablissement_stages
  * etablissement_diplomes

Approche : idempotent (UPSERT / INSERT ON CONFLICT).
Pour chaque établissement XLSX, on récupère son `fiche_id` via
`etablissement_details_xlsx.id_source_xlsx`, puis on remplit :
- les nouvelles colonnes de fiches_etablissement depuis toutes les feuilles
- les tables satellites

Usage :
    python3 scripts/enrichir_etablissement_complet.py --dry-run
    python3 scripts/enrichir_etablissement_complet.py
"""
from __future__ import annotations

import argparse
import subprocess
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "Base_Etablissements_Superieurs_Togo.xlsx"
DB_NAME = "activ_education"
DB_USER = "postgres"
CONTAINER_NAME = "activeducation-db"


def run_psql(sql: str) -> str:
    res = subprocess.run(
        ["docker", "exec", CONTAINER_NAME, "psql", "-U", DB_USER, "-d", DB_NAME,
         "-t", "-A", "-c", sql],
        capture_output=True, text=True, check=True,
    )
    return res.stdout.strip()


def cell_str(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_non_disponible(s: str) -> bool:
    return not s or s.lower() in {"non disponible", "n/a", "nd", "-", "—", ""}


def sql_escape(s: str) -> str:
    return s.replace("'", "''")


# ── Lecture XLSX ─────────────────────────────────────────────────────────────
def charger_xlsx_par_feuille() -> dict[str, list[dict]]:
    """Renvoie {nom_feuille: [{col: val, ...}, ...]}."""
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    out: dict[str, list[dict]] = {}
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        headers = [cell_str(c.value) for c in next(ws.iter_rows(min_row=1, max_row=1))]
        rows = []
        for row in ws.iter_rows(min_row=2, values_only=True):
            if not row or all(cell is None for cell in row):
                continue
            d = {}
            for h, v in zip(headers, row):
                d[h] = v
            rows.append(d)
        out[sheet_name] = rows
    return out


def index_par_id_etablissement(rows: list[dict]) -> dict[str, dict]:
    """Indexe les lignes par ID_ETABLISSEMENT (1 ligne par établissement
    attendu ; en cas de doublons, on garde la première)."""
    out: dict[str, dict] = {}
    for r in rows:
        eid = cell_str(r.get("ID_ETABLISSEMENT"))
        if eid and eid not in out:
            out[eid] = r
    return out


def mapping_db_xlsx() -> dict[str, int]:
    raw = run_psql(
        "SELECT id_source_xlsx, fiche_id FROM etablissement_details_xlsx "
        "WHERE id_source_xlsx IS NOT NULL;"
    )
    out: dict[str, int] = {}
    for line in raw.splitlines():
        if "|" not in line:
            continue
        xid, fid = line.split("|", 1)
        out[xid.strip()] = int(fid.strip())
    return out


# ── Conversions XLSX → DB ────────────────────────────────────────────────────
def parse_bool(v) -> bool | None:
    s = cell_str(v).lower()
    if is_non_disponible(s):
        return None
    if s in {"oui", "yes", "true", "1", "vrai", "disponible", "présent", "present"}:
        return True
    if s in {"non", "no", "false", "0", "faux", "absent", "indisponible"}:
        return False
    return None


def parse_int(v) -> int | None:
    s = cell_str(v)
    if is_non_disponible(s):
        return None
    try:
        return int(s.replace(" ", ""))
    except ValueError:
        return None


def parse_float(v) -> float | None:
    s = cell_str(v).replace(",", ".")
    if is_non_disponible(s):
        return None
    try:
        return float(s)
    except ValueError:
        return None


# ── Mise à jour DB ───────────────────────────────────────────────────────────
def update_etablissement(fiche_id: int, values: dict) -> None:
    """values = {col: valeur} — ne met à jour que les colonnes présentes."""
    if not values:
        return
    sets = []
    for col, val in values.items():
        if val is None:
            continue
        if isinstance(val, bool):
            sets.append(f"{col} = {val}")
        elif isinstance(val, (int, float)):
            sets.append(f"{col} = {val}")
        else:
            sets.append(f"{col} = '{sql_escape(str(val))}'")
    if not sets:
        return
    sql = f"UPDATE fiches_etablissement SET {', '.join(sets)} WHERE id = {fiche_id};"
    run_psql(sql)


def upsert_reseaux_sociaux(fiche_id: int, reseaux: list[tuple[str, str]]) -> None:
    """reseaux = [(type, url), ...]"""
    # Supprime puis insère (idempotent simple)
    run_psql(f"DELETE FROM etablissement_reseaux_sociaux WHERE etablissement_id = {fiche_id};")
    for type_, url in reseaux:
        run_psql(
            f"INSERT INTO etablissement_reseaux_sociaux (etablissement_id, type, url) "
            f"VALUES ({fiche_id}, '{type_}', '{sql_escape(url)}') "
            f"ON CONFLICT DO NOTHING;"
        )


def upsert_admissions(fiche_id: int, a: dict) -> None:
    keys = ["niveau_requis", "serie_bac_acceptee", "concours", "etude_dossier",
            "entretien", "test", "age_minimum", "nombre_places",
            "dates_inscription", "dates_concours"]
    cols = ["etablissement_id"] + keys
    vals = [str(fiche_id)]
    for k in keys:
        v = a.get(k)
        if v is None:
            vals.append("NULL")
        elif isinstance(v, bool):
            vals.append(str(v))
        elif isinstance(v, (int, float)):
            vals.append(str(v))
        else:
            vals.append(f"'{sql_escape(str(v))}'")

    placeholders = ", ".join(["%s"] * len(cols))
    col_list = ", ".join(cols)
    val_list = ", ".join(vals)

    # INSERT ... ON CONFLICT (etablissement_id) DO UPDATE SET ...
    updates = ", ".join(f"{k} = EXCLUDED.{k}" for k in keys)
    sql = (
        f"INSERT INTO etablissement_admissions ({col_list}) VALUES ({val_list}) "
        f"ON CONFLICT (etablissement_id) DO UPDATE SET {updates};"
    )
    run_psql(sql)


def upsert_responsables(fiche_id: int, r: dict) -> None:
    keys = ["directeur", "recteur_doyen", "president", "fondateur",
            "responsable_admissions", "responsable_communication", "contacts_administratifs"]
    cols = ["etablissement_id"] + keys
    vals = [str(fiche_id)]
    for k in keys:
        v = r.get(k)
        if v is None:
            vals.append("NULL")
        else:
            vals.append(f"'{sql_escape(str(v))}'")

    placeholders = ", ".join(["%s"] * len(cols))
    col_list = ", ".join(cols)
    val_list = ", ".join(vals)

    updates = ", ".join(f"{k} = EXCLUDED.{k}" for k in keys)
    sql = (
        f"INSERT INTO etablissement_responsables ({col_list}) VALUES ({val_list}) "
        f"ON CONFLICT (etablissement_id) DO UPDATE SET {updates};"
    )
    run_psql(sql)


def upsert_stages(fiche_id: int, s: dict) -> None:
    keys = ["stage_obligatoire", "duree_stage", "entreprises_partenaires",
            "alternance", "bureau_carriere", "offres_emploi", "reseau_anciens"]
    cols = ["etablissement_id"] + keys
    vals = [str(fiche_id)]
    for k in keys:
        v = s.get(k)
        if v is None:
            vals.append("NULL")
        elif isinstance(v, bool):
            vals.append(str(v))
        else:
            vals.append(f"'{sql_escape(str(v))}'")

    col_list = ", ".join(cols)
    val_list = ", ".join(vals)
    updates = ", ".join(f"{k} = EXCLUDED.{k}" for k in keys)
    sql = (
        f"INSERT INTO etablissement_stages ({col_list}) VALUES ({val_list}) "
        f"ON CONFLICT (etablissement_id) DO UPDATE SET {updates};"
    )
    run_psql(sql)


def upsert_diplomes(fiche_id: int, diplomes: list[str]) -> None:
    run_psql(f"DELETE FROM etablissement_diplomes WHERE etablissement_id = {fiche_id};")
    for d in diplomes:
        run_psql(
            f"INSERT INTO etablissement_diplomes (etablissement_id, diplome) "
            f"VALUES ({fiche_id}, '{sql_escape(d)}') ON CONFLICT DO NOTHING;"
        )


# ── Mapping XLSX → DB ───────────────────────────────────────────────────────
def map_type_etablissement(s: str) -> str | None:
    """Mappe le label XLSX (français) vers l'enum CHECK de la DB."""
    if not s or is_non_disponible(s):
        return None
    s_low = s.lower()
    if "université" in s_low or "universite" in s_low:
        return "UNIVERSITE"
    if "grande école" in s_low or "grande ecole" in s_low:
        return "GRANDE_ECOLE"
    if "lycée" in s_low or "lycee" in s_low:
        return "LYCEE"
    if "collège" in s_low or "college" in s_low:
        return "COLLEGE"
    if "centre" in s_low and ("formation" in s_low or "professionnelle" in s_low):
        return "CENTRE_FORMATION_PROFESSIONNELLE"
    return "AUTRE"


def build_etablissement_update(eid_data: dict) -> dict:
    """Construit le dict {col: valeur} à partir d'une ligne XLSX 'Etablissements'."""
    out = {}
    # Identité
    if not is_non_disponible(eid_data.get("Sigle", "")):
        out["sigle"] = cell_str(eid_data.get("Sigle"))
    if not is_non_disponible(eid_data.get("Logo_URL", "")):
        out["logo_url"] = cell_str(eid_data.get("Logo_URL"))
    if not is_non_disponible(eid_data.get("Devise", "")):
        out["devise"] = cell_str(eid_data.get("Devise"))
    if not is_non_disponible(eid_data.get("Slogan", "")):
        out["slogan"] = cell_str(eid_data.get("Slogan"))
    if not is_non_disponible(eid_data.get("Statut juridique", "")):
        out["statut_juridique"] = cell_str(eid_data.get("Statut juridique"))
    if not is_non_disponible(eid_data.get("Numéro d'agrément", "")):
        out["numero_agrement"] = cell_str(eid_data.get("Numéro d'agrément"))
    if not is_non_disponible(eid_data.get("Date d'agrément", "")):
        out["date_agrement"] = cell_str(eid_data.get("Date d'agrément"))
    if not is_non_disponible(eid_data.get("Autorité de tutelle", "")):
        out["autorite_tutelle"] = cell_str(eid_data.get("Autorité de tutelle"))
    if not is_non_disponible(eid_data.get("Accréditations", "")):
        out["accreditations"] = cell_str(eid_data.get("Accréditations"))
    rc = parse_bool(eid_data.get("Reconnaissance CAMES"))
    if rc is not None:
        out["reconnaissance_cames"] = rc
    y = parse_int(eid_data.get("Année de création"))
    if y is not None:
        out["annee_creation"] = y
    for src, dst in [
        ("Historique", "historique"),
        ("Présentation", "presentation"),
        ("Mission", "mission"),
        ("Vision", "vision"),
        ("Valeurs", "valeurs"),
    ]:
        v = cell_str(eid_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    # Type d'établissement : mapping vers l'enum CHECK de la DB
    type_xlsx = eid_data.get("Type", "")
    type_db = map_type_etablissement(type_xlsx)
    if type_db is not None:
        out["type_etablissement"] = type_db
    return out


def build_localisation_update(loc_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Région", "region"),
        ("Préfecture", "prefecture"),
        ("Commune", "commune"),
        ("Quartier", "quartier"),
    ]:
        v = cell_str(loc_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    if not is_non_disponible(loc_data.get("Lien Google Maps", "")):
        out["lien_google_maps"] = cell_str(loc_data.get("Lien Google Maps"))
    if not is_non_disponible(loc_data.get("Plan du campus", "")):
        out["plan_campus"] = cell_str(loc_data.get("Plan du campus"))
    return out


def build_calendrier_update(cal_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Ouverture des inscriptions", "ouverture_inscriptions"),
        ("Clôture", "cloture_inscriptions"),
        ("Dates des concours", "dates_concours"),
        ("Rentrée", "date_rentree"),
        ("Vacances", "vacances"),
        ("Examens", "examens"),
        ("Soutenances", "soutenances"),
    ]:
        v = cell_str(cal_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    return out


def build_frais_update(frais_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Frais d'inscription", "frais_inscription"),
        ("Réinscription", "reinscription"),
        ("Scolarité annuelle", "scolarite_annuelle"),
        ("Mensualités", "mensualites"),
        ("Assurance", "frais_assurance"),
        ("Uniforme", "frais_uniforme"),
        ("Carte étudiant", "frais_carte_etudiant"),
        ("Frais de laboratoire", "frais_laboratoire"),
        ("Examens", "frais_examens"),
        ("Bibliothèque", "frais_bibliotheque"),
        ("Autres frais", "frais_autres"),
        ("Devise", "frais_devise"),
    ]:
        v = cell_str(frais_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    pe = parse_bool(frais_data.get("Paiement échelonné"))
    if pe is not None:
        out["paiement_echelonne"] = pe
    return out


def build_infra_update(infra_data: dict) -> dict:
    out = {}
    if not is_non_disponible(infra_data.get("Synthèse infrastructures", "")):
        out["infrastructures_synthese"] = cell_str(infra_data.get("Synthèse infrastructures"))
    for src, dst in [
        ("Nombre de bâtiments", "nb_batiments"),
        ("Amphithéâtres", "nb_amphitheatres"),
        ("Salles de classe", "nb_salles_classe"),
        ("Laboratoires", "nb_laboratoires"),
    ]:
        v = parse_int(infra_data.get(src, ""))
        if v is not None:
            out[dst] = v
    # Booléens depuis colonnes directes (rarement remplis)
    for src, dst in [
        ("Bibliothèque", "bibliotheque"),
        ("Bibliothèque numérique", "bibliotheque_numerique"),
        ("Salle informatique", "salle_informatique"),
        ("FabLab", "fablab"),
        ("Incubateur", "incubateur"),
        ("Wi-Fi", "wifi"),
        ("Internat", "internat"),
        ("Restaurant", "restaurant"),
        ("Cafétéria", "cafeteria"),
        ("Terrain de sport", "terrain_sport"),
        ("Infirmerie", "infirmerie"),
        ("Parking", "parking"),
        ("Accessibilité PMR", "accessibilite_pmr"),
        ("Sécurité", "securite"),
    ]:
        v = parse_bool(infra_data.get(src, ""))
        if v is not None:
            out[dst] = v

    # Heuristique sur la synthèse (souvent seule source d'info)
    synthese = cell_str(infra_data.get("Synthèse infrastructures", "")).lower()
    if synthese and not is_non_disponible(synthese):
        keyword_map = {
            "bibliothèque": "bibliotheque",
            "bibliotheque": "bibliotheque",
            "bibliothèque numérique": "bibliotheque_numerique",
            "salle informatique": "salle_informatique",
            "laboratoire": "salle_informatique",  # approximation
            "fablab": "fablab",
            "incubateur": "incubateur",
            "wi-fi": "wifi",
            "wifi": "wifi",
            "cité universitaire": "internat",
            "internat": "internat",
            "restaurant universitaire": "restaurant",
            "restaurant": "restaurant",
            "cafétéria": "cafeteria",
            "cafeteria": "cafeteria",
            "terrain de sport": "terrain_sport",
            "terrain sportif": "terrain_sport",
            "infirmerie": "infirmerie",
            "parking": "parking",
            "pmr": "accessibilite_pmr",
            "sécurité": "securite",
            "securite": "securite",
        }
        for kw, dst in keyword_map.items():
            if kw in synthese and dst not in out:
                out[dst] = True
    return out


def build_vie_etudiante_update(ve_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Clubs", "clubs"),
        ("Associations", "associations"),
        ("Activités culturelles", "activites_culturelles"),
        ("Sports", "sports"),
        ("Compétitions", "competitions"),
        ("Conférences", "conferences"),
        ("Événements", "evenements"),
    ]:
        v = cell_str(ve_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    be = parse_bool(ve_data.get("Bureau des étudiants"))
    if be is not None:
        out["bureau_etudiants"] = be
    ha = parse_bool(ve_data.get("Hackathons"))
    if ha is not None:
        out["hackathons"] = ha
    return out


def build_enseignants_update(en_data: dict) -> dict:
    out = {}
    if not is_non_disponible(en_data.get("Corps enseignant (synthèse)", "")):
        out["corps_enseignant_synthese"] = cell_str(en_data.get("Corps enseignant (synthèse)"))
    for src, dst in [
        ("Nombre d'enseignants", "nb_enseignants"),
        ("Permanents", "enseignants_permanents"),
        ("Vacataires", "enseignants_vacataires"),
        ("Professeurs titulaires", "enseignants_titulaires"),
        ("Enseignants docteurs", "enseignants_docteurs"),
    ]:
        v = parse_int(en_data.get(src, ""))
        if v is not None:
            out[dst] = v
    ep = parse_bool(en_data.get("Experts professionnels"))
    if ep is not None:
        out["enseignants_experts_pro"] = ep
    return out


def build_stats_update(st_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Nombre d'étudiants", "nb_etudiants"),
        ("Nombre de diplômés", "nb_diplomes"),
    ]:
        v = parse_int(st_data.get(src, ""))
        if v is not None:
            out[dst] = v
    for src, dst in [
        ("Taux de réussite", "taux_reussite"),
        ("Taux d'abandon", "taux_abandon"),
        ("Ratio enseignants/étudiants", "ratio_enseignants_etudiants"),
        ("Taux d'insertion professionnelle", "taux_insertion_pro"),
        ("Satisfaction", "satisfaction"),
    ]:
        v = cell_str(st_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    return out


def build_partenariats_update(pa_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Partenaires (synthèse)", "partenaires_synthese"),
        ("Universités partenaires", "universites_partenaires"),
        ("Entreprises partenaires", "entreprises_partenaires"),
        ("ONG", "ong_partenaires"),
        ("Ministères", "ministeres_partenaires"),
        ("Organisations internationales", "org_internationales_partenaires"),
    ]:
        v = cell_str(pa_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    for src, dst in [
        ("Doubles diplômes", "doubles_diplomes"),
        ("Programmes de mobilité", "programmes_mobilite"),
    ]:
        v = parse_bool(pa_data.get(src, ""))
        if v is not None:
            out[dst] = v
    return out


def build_recherche_update(re_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Recherche (synthèse)", "recherche_synthese"),
        ("Laboratoires", "laboratoires"),
        ("Publications", "publications"),
        ("Projets", "projets_recherche"),
        ("Brevets", "brevets"),
        ("Innovations", "innovations"),
    ]:
        v = cell_str(re_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    return out


def build_bourses_update(bo_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Bourses internes", "bourses_internes"),
        ("Bourses gouvernementales", "bourses_gouvernementales"),
        ("Bourses internationales", "bourses_internationales"),
        ("Critères", "bourses_criteres"),
        ("Montants", "bourses_montants"),
    ]:
        v = cell_str(bo_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    return out


def build_documents_update(do_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Brochure", "brochure_url"),
        ("Guide étudiant", "guide_etudiant_url"),
        ("Règlement intérieur", "reglement_interieur_url"),
        ("Calendrier académique", "calendrier_academique_url"),
        ("Dossier d'inscription", "dossier_inscription_url"),
        ("Formulaire d'admission", "formulaire_admission_url"),
    ]:
        v = cell_str(do_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    return out


def build_avis_update(av_data: dict) -> dict:
    out = {}
    nm = parse_float(av_data.get("Note moyenne"))
    if nm is not None:
        out["note_moyenne_avis"] = nm
    na = parse_int(av_data.get("Nombre d'avis"))
    if na is not None:
        out["nb_avis"] = na
    for src, dst in [
        ("Résumé des avis positifs", "avis_positifs_resume"),
        ("Résumé des critiques récurrentes", "avis_critiques_resume"),
    ]:
        v = cell_str(av_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    return out


def extract_reseaux_sociaux(co_data: dict) -> list[tuple[str, str]]:
    """Renvoie [(type, url), ...] depuis la feuille Contacts."""
    out = []
    mapping = [
        ("WhatsApp", "whatsapp"),
        ("Facebook", "facebook"),
        ("LinkedIn", "linkedin"),
        ("Instagram", "instagram"),
        ("TikTok", "tiktok"),
        ("YouTube", "youtube"),
        ("Telegram", "telegram"),
    ]
    for src, type_ in mapping:
        v = cell_str(co_data.get(src, ""))
        if not is_non_disponible(v):
            out.append((type_, v))
    return out


def extract_admissions(ad_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Niveau requis", "niveau_requis"),
        ("Série du BAC acceptée", "serie_bac_acceptee"),
        ("Dates d'inscription", "dates_inscription"),
        ("Dates de concours", "dates_concours"),
    ]:
        v = cell_str(ad_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    for src, dst in [
        ("Concours", "concours"),
        ("Étude de dossier", "etude_dossier"),
        ("Entretien", "entretien"),
        ("Test", "test"),
    ]:
        v = parse_bool(ad_data.get(src, ""))
        if v is not None:
            out[dst] = v
    am = parse_int(ad_data.get("Âge minimum"))
    if am is not None:
        out["age_minimum"] = am
    np = parse_int(ad_data.get("Nombre de places"))
    if np is not None:
        out["nombre_places"] = np
    return out


def extract_responsables(re_data: dict) -> dict:
    out = {}
    mapping = [
        ("Directeur", "directeur"),
        ("Recteur / Directeur / Doyen", "recteur_doyen"),
        ("Président", "president"),
        ("Fondateur", "fondateur"),
        ("Responsable admissions", "responsable_admissions"),
        ("Responsable communication", "responsable_communication"),
        ("Contacts administratifs", "contacts_administratifs"),
    ]
    for src, dst in mapping:
        v = cell_str(re_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    return out


def extract_stages(st_data: dict) -> dict:
    out = {}
    for src, dst in [
        ("Durée", "duree_stage"),
        ("Entreprises partenaires", "entreprises_partenaires"),
        ("Offres d'emploi", "offres_emploi"),
        ("Réseau des anciens étudiants", "reseau_anciens"),
    ]:
        v = cell_str(st_data.get(src, ""))
        if not is_non_disponible(v):
            out[dst] = v
    for src, dst in [
        ("Stage obligatoire", "stage_obligatoire"),
        ("Alternance", "alternance"),
        ("Bureau carrière", "bureau_carriere"),
    ]:
        v = parse_bool(st_data.get(src, ""))
        if v is not None:
            out[dst] = v
    return out


def extract_diplomes(di_data: dict) -> list[str]:
    """Renvoie la liste des diplômes où la valeur XLSX ≠ 'Non disponible'."""
    out = []
    mapping = [
        "CAP", "CQP", "BT", "BTS", "DUT",
        "Licence", "Licence Professionnelle",
        "Master", "Master Professionnel", "Doctorat", "Certifications",
    ]
    for d in mapping:
        v = cell_str(di_data.get(d, ""))
        if not is_non_disponible(v):
            out.append(d)
    return out


# ── Main ─────────────────────────────────────────────────────────────────────
def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    args = parser.parse_args()

    print("Lecture du XLSX (toutes feuilles)...")
    data = charger_xlsx_par_feuille()
    print(f"  → {len(data)} feuilles lues")

    # Index par ID_ETABLISSEMENT
    idx = {name: index_par_id_etablissement(rows) for name, rows in data.items()}

    print("Lecture du mapping DB ↔ XLSX...")
    mapping = mapping_db_xlsx()
    print(f"  → {len(mapping)} établissements à enrichir")

    counters = {
        "etablissements_updated": 0,
        "reseaux": 0,
        "admissions": 0,
        "responsables": 0,
        "stages": 0,
        "diplomes": 0,
    }

    for xlsx_id, fiche_id in mapping.items():
        # 1. Update de la grosse fiche
        all_updates = {}
        for builder, sheet_name in [
            (build_etablissement_update, "Etablissements"),
            (build_localisation_update, "Localisation"),
            (build_calendrier_update, "Calendrier"),
            (build_frais_update, "Frais"),
            (build_infra_update, "Infrastructures"),
            (build_vie_etudiante_update, "Vie etudiante"),
            (build_enseignants_update, "Enseignants"),
            (build_stats_update, "Statistiques"),
            (build_partenariats_update, "Partenariats"),
            (build_recherche_update, "Recherche"),
            (build_bourses_update, "Bourses"),
            (build_documents_update, "Documents"),
            (build_avis_update, "Avis"),
        ]:
            sheet_data = idx.get(sheet_name, {}).get(xlsx_id, {})
            if sheet_data:
                all_updates.update(builder(sheet_data))

        if all_updates:
            if not args.dry_run:
                update_etablissement(fiche_id, all_updates)
            counters["etablissements_updated"] += 1

        # 2. Réseaux sociaux
        reseaux = extract_reseaux_sociaux(idx.get("Contacts", {}).get(xlsx_id, {}))
        if reseaux:
            if not args.dry_run:
                upsert_reseaux_sociaux(fiche_id, reseaux)
            counters["reseaux"] += len(reseaux)

        # 3. Admissions
        admissions = extract_admissions(idx.get("Admissions", {}).get(xlsx_id, {}))
        if admissions:
            if not args.dry_run:
                upsert_admissions(fiche_id, admissions)
            counters["admissions"] += 1

        # 4. Responsables
        resp = extract_responsables(idx.get("Responsables", {}).get(xlsx_id, {}))
        if resp:
            if not args.dry_run:
                upsert_responsables(fiche_id, resp)
            counters["responsables"] += 1

        # 5. Stages
        stages = extract_stages(idx.get("Stages", {}).get(xlsx_id, {}))
        if stages:
            if not args.dry_run:
                upsert_stages(fiche_id, stages)
            counters["stages"] += 1

        # 6. Diplômes
        diplomes = extract_diplomes(idx.get("Diplomes", {}).get(xlsx_id, {}))
        if diplomes:
            if not args.dry_run:
                upsert_diplomes(fiche_id, diplomes)
            counters["diplomes"] += len(diplomes)

    # Rapport
    print("\n" + "=" * 70)
    print(f"RÉSULTAT  (mode {'DRY-RUN' if args.dry_run else 'UPDATE'})")
    print("=" * 70)
    for k, v in counters.items():
        print(f"  {k:30} : {v}")


if __name__ == "__main__":
    main()
