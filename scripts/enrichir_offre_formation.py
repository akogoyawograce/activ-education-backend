#!/usr/bin/env python3
"""
Enrichit `fiches_etablissement.offre_formation` à partir de la feuille
`Formations` du XLSX `Base_Etablissements_Superieurs_Togo.xlsx`.

Le lien DB ↔ XLSX passe par la table `etablissement_details_xlsx.id_source_xlsx`
(= l'`ID_ETABLISSEMENT` du XLSX, ex. `TG-UL-001`).

Stratégie :
- Agrège les formations d'un établissement en texte structuré lisible,
  regroupées par Domaine, avec Filière / Niveau / Diplôme.
- N'écrase que les fiches dont l'offre est NULL, vide ou vaut 'Non disponible'
  (les admins peuvent avoir renseigné du contenu manuellement).

Usage :
    python3 scripts/enrichir_offre_formation.py
    python3 scripts/enrichir_offre_formation.py --dry-run   # Aperçu sans UPDATE
"""
from __future__ import annotations

import argparse
import os
import subprocess
import sys
from collections import defaultdict
from pathlib import Path

try:
    import openpyxl
except ImportError:
    print("openpyxl manquant. Installer : pip install openpyxl", file=sys.stderr)
    sys.exit(1)

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "Base_Etablissements_Superieurs_Togo.xlsx"
DB_NAME = "activ_education"
DB_USER = "postgres"
CONTAINER_NAME = "activeducation-db"


def run_psql(sql: str) -> str:
    """Exécute une requête SQL via psql dans le conteneur Docker et renvoie stdout."""
    result = subprocess.run(
        ["docker", "exec", CONTAINER_NAME, "psql", "-U", DB_USER, "-d", DB_NAME,
         "-t", "-A", "-c", sql],
        capture_output=True, text=True, check=True,
    )
    return result.stdout.strip()


def charger_formations_xlsx() -> dict[str, list[dict]]:
    """Renvoie {ID_ETABLISSEMENT: [{domaine, filiere, option, niveau, diplome}, ...]}."""
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Formations"]
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    headers = [str(c).strip() if c else "" for c in next(
        openpyxl.load_workbook(XLSX_PATH, data_only=True)["Formations"].iter_rows(min_row=1, max_row=1, values_only=True)
    )]
    idx = {h: i for i, h in enumerate(headers)}

    out: dict[str, list[dict]] = defaultdict(list)
    for row in rows:
        if not row or not row[idx["ID_ETABLISSEMENT"]]:
            continue
        eid = str(row[idx["ID_ETABLISSEMENT"]]).strip()
        out[eid].append({
            "domaine": (str(row[idx["Domaine"]]).strip()
                        if row[idx["Domaine"]] else ""),
            "filiere": (str(row[idx["Filière"]]).strip()
                        if row[idx["Filière"]] else ""),
            "option": (str(row[idx["Option"]]).strip()
                       if row[idx["Option"]] else ""),
            "niveau": (str(row[idx["Niveau"]]).strip()
                       if row[idx["Niveau"]] else ""),
            "diplome": (str(row[idx["Diplôme préparé"]]).strip()
                        if row[idx["Diplôme préparé"]] else ""),
        })
    return out


def formater_offre(formations: list[dict]) -> str:
    """Texte structuré lisible, regroupé par domaine."""
    # Bucket par domaine, en gardant l'ordre de première apparition
    by_domaine: dict[str, list[dict]] = {}
    for f in formations:
        d = f["domaine"] or "Autre"
        by_domaine.setdefault(d, []).append(f)

    lines: list[str] = []
    for domaine, items in by_domaine.items():
        lines.append(f"• {domaine}")
        for it in items:
            sub_parts = [it["filiere"]] if it["filiere"] else []
            if it["option"] and it["option"].lower() != "non disponible":
                sub_parts.append(f"option {it['option']}")
            if it["niveau"] and it["niveau"].lower() != "non disponible":
                sub_parts.append(it["niveau"])
            if it["diplome"] and it["diplome"].lower() != "non disponible":
                sub_parts.append(it["diplome"])
            lines.append(f"   – { ' — '.join(sub_parts) }")
    return "\n".join(lines)


def mapping_db_xlsx() -> dict[str, int]:
    """Renvoie {id_source_xlsx: fiche_id} depuis la DB."""
    out_raw = run_psql(
        "SELECT id_source_xlsx, fiche_id FROM etablissement_details_xlsx "
        "WHERE id_source_xlsx IS NOT NULL;"
    )
    out: dict[str, int] = {}
    for line in out_raw.splitlines():
        if "|" not in line:
            continue
        xlsx_id, fiche_id = line.split("|", 1)
        out[xlsx_id.strip()] = int(fiche_id.strip())
    return out


def fiche_actuelle(fiche_id: int) -> str | None:
    val = run_psql(
        f"SELECT offre_formation FROM fiches_etablissement WHERE id = {fiche_id};"
    )
    if val == "":
        return None
    return val


def update_offre(fiche_id: int, contenu: str) -> None:
    # Échapper les apostrophes pour SQL
    contenu_escaped = contenu.replace("'", "''")
    run_psql(
        f"UPDATE fiches_etablissement "
        f"SET offre_formation = '{contenu_escaped}' "
        f"WHERE id = {fiche_id};"
    )


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true",
                        help="Aperçu sans modification de la DB")
    parser.add_argument("--limit", type=int, default=0,
                        help="Limiter le nombre de fiches traitées (0 = toutes)")
    args = parser.parse_args()

    if not XLSX_PATH.exists():
        print(f"XLSX introuvable : {XLSX_PATH}", file=sys.stderr)
        sys.exit(1)

    print(f"Lecture du XLSX : {XLSX_PATH.name}")
    formations_par_etab = charger_formations_xlsx()
    print(f"  → {len(formations_par_etab)} établissements avec formations")

    print("Lecture du mapping DB ↔ XLSX...")
    mapping = mapping_db_xlsx()
    print(f"  → {len(mapping)} fiches liées au XLSX")

    updated = 0
    skipped_non_vide = 0
    skipped_pas_de_formation = 0
    apercu: list[tuple[int, str, str]] = []  # (fiche_id, titre, nouveau_contenu)

    for xlsx_id, fiche_id in mapping.items():
        forms = formations_par_etab.get(xlsx_id, [])
        if not forms:
            skipped_pas_de_formation += 1
            continue

        # Récupère le titre de la fiche pour l'aperçu
        titre = run_psql(
            f"SELECT titre FROM fiches WHERE id = {fiche_id};"
        ) or f"(fiche {fiche_id})"

        actuel = fiche_actuelle(fiche_id)
        # Ne pas écraser si déjà renseigné
        if actuel and actuel.strip() and actuel.strip().lower() != "non disponible":
            skipped_non_vide += 1
            continue

        contenu = formater_offre(forms)
        apercu.append((fiche_id, titre, contenu))

        if not args.dry_run:
            update_offre(fiche_id, contenu)
            updated += 1

        if args.limit and updated >= args.limit:
            break

    # Rapport
    print("\n" + "=" * 70)
    print(f"RÉSULTAT  (mode {'DRY-RUN' if args.dry_run else 'UPDATE'})")
    print("=" * 70)
    print(f"Fiches éligibles (offre vide)        : {len(apercu)}")
    print(f"Fiches mises à jour                  : {updated}")
    print(f"Skipped — offre déjà renseignée      : {skipped_non_vide}")
    print(f"Skipped — pas de formations XLSX     : {skipped_pas_de_formation}")

    if apercu:
        print("\n--- APERÇU (3 premiers) ---")
        for fid, titre, contenu in apercu[:3]:
            print(f"\n[Fiche #{fid}] {titre}")
            # Tronque l'aperçu
            lines = contenu.split("\n")
            for ln in lines[:8]:
                print(f"  {ln}")
            if len(lines) > 8:
                print(f"  ... (+{len(lines) - 8} lignes)")


if __name__ == "__main__":
    main()
