#!/usr/bin/env python3
"""
Deuxième passe d'enrichissement des liens établissement↔filière :

1. Match strict (exact après normalisation) XLSX → DB pour les fiches existantes
2. Création des fiches filières manquantes les plus fréquentes (≥ 3 établissements)
   + insertion des liens correspondants

IMPORTANT : cette version utilise UNIQUEMENT le matching exact, pour éviter
les faux positifs observés en première passe (ex. 'Agroéconomie' → 'Économie').
"""
from __future__ import annotations

import argparse
import re
import subprocess
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
XLSX_PATH = ROOT / "Base_Etablissements_Superieurs_Togo.xlsx"
DB_NAME = "activ_education"
DB_USER = "postgres"
CONTAINER_NAME = "activeducation-db"

# Domaines prédéfinis (basés sur la liste actuelle en base)
DOMAINES = {
    "agriculture & environnement": "Agriculture & Environnement",
    "agriculture": "Agriculture & Environnement",
    "environnement": "Agriculture & Environnement",
    "agro": "Agriculture & Environnement",
    "agronomie": "Agriculture & Environnement",
    "forest": "Agriculture & Environnement",
    "lettres & sciences humaines": "Lettres & Sciences Humaines",
    "lettres": "Lettres & Sciences Humaines",
    "sciences humaines": "Lettres & Sciences Humaines",
    "anglais": "Lettres & Sciences Humaines",
    "francais": "Lettres & Sciences Humaines",
    "allemand": "Lettres & Sciences Humaines",
    "arabe": "Lettres & Sciences Humaines",
    "espagnol": "Lettres & Sciences Humaines",
    "histoire": "Lettres & Sciences Humaines",
    "geographie": "Lettres & Sciences Humaines",
    "sociologie": "Lettres & Sciences Humaines",
    "psychologie": "Lettres & Sciences Humaines",
    "philosophie": "Lettres & Sciences Humaines",
    "anthropologie": "Lettres & Sciences Humaines",
    "éducation": "Éducation",
    "staps": "Éducation",
    "sport": "Éducation",
    "enseignement": "Éducation",
    "sciences & technologies": "Sciences & Technologies",
    "sciences": "Sciences & Technologies",
    "technologie": "Sciences & Technologies",
    "technologies": "Sciences & Technologies",
    "genie": "Sciences & Technologies",
    "informatique": "Sciences & Technologies",
    "mathematiques": "Sciences & Technologies",
    "physique": "Sciences & Technologies",
    "chimie": "Sciences & Technologies",
    "biologie": "Sciences & Technologies",
    "telecommunication": "Sciences & Technologies",
    "cybersecurite": "Sciences & Technologies",
    "architecture": "Sciences & Technologies",
    "économie & gestion": "Économie & Gestion",
    "économie": "Économie & Gestion",
    "economie": "Économie & Gestion",
    "gestion": "Économie & Gestion",
    "management": "Économie & Gestion",
    "comptabilite": "Économie & Gestion",
    "finance": "Économie & Gestion",
    "banque": "Économie & Gestion",
    "assurance": "Économie & Gestion",
    "commerce": "Économie & Gestion",
    "marketing": "Économie & Gestion",
    "transport": "Économie & Gestion",
    "logistique": "Économie & Gestion",
    "droit & politique": "Droit & Politique",
    "droit": "Droit & Politique",
    "science politique": "Droit & Politique",
    "administration publique": "Droit & Politique",
    "médecine & santé": "Médecine & Santé",
    "médecine": "Médecine & Santé",
    "medecine": "Médecine & Santé",
    "pharmacie": "Médecine & Santé",
    "sante": "Médecine & Santé",
    "infirmier": "Médecine & Santé",
    "sage-femme": "Médecine & Santé",
    "arts & culture": "Arts & Culture",
    "arts": "Arts & Culture",
    "culture": "Arts & Culture",
    "communication": "Arts & Culture",
    "journalisme": "Arts & Culture",
    "tourisme": "Arts & Culture",
    "hotellerie": "Arts & Culture",
    "audiovisuel": "Arts & Culture",
    "design": "Arts & Culture",
}


def run_psql(sql: str) -> str:
    res = subprocess.run(
        ["docker", "exec", CONTAINER_NAME, "psql", "-U", DB_USER, "-d", DB_NAME,
         "-t", "-A", "-c", sql],
        capture_output=True, text=True, check=True,
    )
    return res.stdout.strip()


def normalize(s: str) -> str:
    if not s:
        return ""
    s = unicodedata.normalize("NFD", s)
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    s = s.lower().strip()
    s = " ".join(s.split())
    return s


def sql_escape(s: str) -> str:
    return s.replace("'", "''")


# ── Lecture XLSX ─────────────────────────────────────────────────────────────
def charger_xlsx_relations() -> dict[str, set[str]]:
    wb = openpyxl.load_workbook(XLSX_PATH, data_only=True)
    ws = wb["Formations"]
    headers = [str(c.value).strip() if c.value else "" for c in next(
        ws.iter_rows(min_row=1, max_row=1))]
    idx = {h: i for i, h in enumerate(headers)}
    out: dict[str, set[str]] = defaultdict(set)
    for row in ws.iter_rows(min_row=2, values_only=True):
        if not row or not row[idx["ID_ETABLISSEMENT"]]:
            continue
        eid = str(row[idx["ID_ETABLISSEMENT"]]).strip()
        filiere = str(row[idx["Filière"]]).strip() if row[idx["Filière"]] else ""
        if filiere:
            out[eid].add(filiere)
    return out


def charger_filieres_db() -> dict[str, int]:
    raw = run_psql(
        "SELECT fi.id, f.titre FROM fiches_filiere fi "
        "JOIN fiches f ON f.id=fi.id "
        "WHERE f.est_publie=true "
        "AND f.titre NOT ILIKE '%test%' AND f.titre NOT ILIKE '%save button%';"
    )
    out: dict[str, int] = {}
    for line in raw.splitlines():
        if "|" not in line:
            continue
        fid, titre = line.split("|", 1)
        out[normalize(titre)] = int(fid.strip())
    return out


def mapping_db_xlsx() -> dict[str, int]:
    raw = run_psql(
        "SELECT id_source_xlsx, fiche_id FROM etablissement_details_xlsx "
        "WHERE id_source_xlsx IS NOT NULL;"
    )
    out: dict[str, int] = {}
    for line in raw.splitlines():
        if "|" not in line:
            continue
        xid, fid = line.split("|", 1)
        out[xid.strip()] = int(fid.strip())
    return out


def lien_existe(etab_id: int, filiere_id: int) -> bool:
    raw = run_psql(
        f"SELECT 1 FROM etablissement_filiere "
        f"WHERE etablissement_id={etab_id} AND filiere_id={filiere_id};"
    )
    return bool(raw.strip())


def inserer_lien(etab_id: int, filiere_id: int) -> None:
    run_psql(
        f"INSERT INTO etablissement_filiere (etablissement_id, filiere_id) "
        f"VALUES ({etab_id}, {filiere_id}) ON CONFLICT DO NOTHING;"
    )


def determiner_domaine(filiere: str) -> str:
    """Devine le domaine à partir du nom de la filière."""
    norm = normalize(filiere)
    for key, val in DOMAINES.items():
        if key in norm:
            return val
    return ""


def determiner_niveau(filiere: str) -> str:
    """Détecte niveau requis à partir du nom."""
    norm = normalize(filiere)
    if "licence pro" in norm or "lpro" in norm:
        return "Bac+2"
    if "master" in norm or "msc" in norm:
        return "Bac+3 / Bac+4"
    if "doctorat" in norm or "phd" in norm:
        return "Bac+5"
    if "bts" in norm or "dut" in norm or "deust" in norm or "bt " in norm or norm.startswith("bt"):
        return "Bac"
    if "ingenieur" in norm or "génie" in norm or "genie" in norm:
        return "Bac C/D/E"
    if "licence" in norm:
        return "Bac"
    return "Bac"


def determiner_duree(filiere: str) -> str:
    norm = normalize(filiere)
    if "doctorat" in norm:
        return "3-5 ans"
    if "master" in norm:
        return "2 ans"
    if "licence pro" in norm or "bts" in norm or "dut" in norm or "deust" in norm:
        return "2-3 ans"
    if "ingenieur" in norm or "génie" in norm or "genie" in norm:
        return "5 ans"
    if "licence" in norm:
        return "3 ans"
    return "3-5 ans"


def creer_fiche_filiere(titre: str, domaine: str, niveau: str, duree: str) -> int:
    """Crée une fiche (titre) + fiche_filiere (domaine, niveau_requis, duree)
    en une transaction. Renvoie l'ID de la fiche filière."""
    titre_esc = sql_escape(titre)
    domaine_esc = sql_escape(domaine) if domaine else ""
    niveau_esc = sql_escape(niveau)
    duree_esc = sql_escape(duree)

    # Crée d'abord la fiche parente en récupérant l'ID via RETURNING
    domaine_sql = f"'{domaine_esc}'" if domaine_esc else "NULL"
    fiche_id_raw = run_psql(
        f"WITH new_fiche AS ("
        f"  INSERT INTO fiches (titre, resume, est_publie, nb_consultations, "
        f"    tracking_id, created_at, updated_at) "
        f"  VALUES ('{titre_esc}', 'Fiche enrichie depuis la base nationale XLSX', "
        f"    true, 0, gen_random_uuid(), NOW(), NOW()) "
        f"  RETURNING id"
        f") "
        f"INSERT INTO fiches_filiere (id, domaine, niveau_requis, duree) "
        f"SELECT id, {domaine_sql}, '{niveau_esc}', '{duree_esc}' FROM new_fiche "
        f"RETURNING id;"
    )
    # psql ajoute 'INSERT 0 1' après le résultat ; on prend la 1ère ligne
    return int(fiche_id_raw.split("\n")[0])


# ── Main ─────────────────────────────────────────────────────────────────────
def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--dry-run", action="store_true")
    parser.add_argument("--min-freq", type=int, default=3,
                        help="Fréquence XLSX minimale pour auto-créer une filière "
                             "manquante (défaut 3)")
    parser.add_argument("--max-creations", type=int, default=30,
                        help="Nombre max de fiches filières à créer (défaut 30)")
    args = parser.parse_args()

    print("Lecture XLSX (Formations)...")
    xlsx_relations = charger_xlsx_relations()

    print("Lecture fiches filières en base...")
    filieres_db = charger_filieres_db()
    print(f"  → {len(filieres_db)} fiches")

    print("Lecture mapping DB ↔ XLSX...")
    mapping = mapping_db_xlsx()

    # ── Étape 1 : matching strict des relations existantes ──────────────────
    matched_total = 0
    inserted_strict = 0

    for xlsx_id, fiche_etab_id in mapping.items():
        for filiere_xlsx in xlsx_relations.get(xlsx_id, set()):
            key = normalize(filiere_xlsx)
            if key in filieres_db:
                matched_total += 1
                fid = filieres_db[key]
                if not lien_existe(fiche_etab_id, fid):
                    if not args.dry_run:
                        inserer_lien(fiche_etab_id, fid)
                        inserted_strict += 1

    print(f"\n[1] Matching strict : {matched_total} relations")
    print(f"    → liens insérés : {inserted_strict}")

    # ── Étape 2 : identifier les filières XLSX non matchées (avec fréquence) ─
    freq_non_matchees: dict[str, set[str]] = defaultdict(set)  # filière → etabs
    for xlsx_id, fiche_etab_id in mapping.items():
        for filiere_xlsx in xlsx_relations.get(xlsx_id, set()):
            key = normalize(filiere_xlsx)
            if key not in filieres_db:
                freq_non_matchees[filiere_xlsx].add(xlsx_id)

    candidats_creation = sorted(
        freq_non_matchees.items(),
        key=lambda x: -len(x[1])
    )
    candidats_creation = [
        (titre, etabs) for titre, etabs in candidats_creation
        if len(etabs) >= args.min_freq
    ][:args.max_creations]

    print(f"\n[2] Filières manquantes (fréquence ≥ {args.min_freq}, "
          f"max {args.max_creations})")
    print(f"    → {len(candidats_creation)} à créer")

    created_count = 0
    new_liens = 0
    for titre, etabs in candidats_creation:
        domaine = determiner_domaine(titre)
        niveau = determiner_niveau(titre)
        duree = determiner_duree(titre)

        if args.dry_run:
            print(f"    [DRY] '{titre}' (dom={domaine}, niv={niveau}, "
                  f"durée={duree}) → {len(etabs)} liens")
            continue

        new_id = creer_fiche_filiere(titre, domaine, niveau, duree)
        # Met à jour la map locale
        filieres_db[normalize(titre)] = new_id
        created_count += 1

        # Crée les liens vers les établissements XLSX
        for xlsx_id in etabs:
            fiche_etab_id = mapping[xlsx_id]
            if not lien_existe(fiche_etab_id, new_id):
                inserer_lien(fiche_etab_id, new_id)
                new_liens += 1

    if not args.dry_run:
        print(f"    → {created_count} fiches créées, {new_liens} liens ajoutés")

    # ── Rapport final ──────────────────────────────────────────────────────
    total_liens = int(run_psql("SELECT COUNT(*) FROM etablissement_filiere;"))
    print("\n" + "=" * 70)
    print(f"RÉSULTAT FINAL")
    print("=" * 70)
    print(f"Total liens etablissement_filiere : {total_liens}")


if __name__ == "__main__":
    main()
