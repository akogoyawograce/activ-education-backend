#!/usr/bin/env python3
"""
Lie les fiches_filiere aux fiches_metier en se basant sur :
  1) Les établissements liés à la filière (via etablissement_filiere)
  2) Les formations XLSX de ces établissements (via etablissement_details_xlsx)
  3) Le secteur d'activité de la formation (feuille 'Debouches')

Stratégie :
  - Récupère la liste des secteurs des formations par établissement.
  - Pour chaque filière, agrège les secteurs de ses établissements liés.
  - Matche chaque filière à TOUTES les fiches_metier dont le secteur matche.

Idempotent : INSERT ... ON CONFLICT DO NOTHING.
"""

import sys
import psycopg2
from collections import defaultdict
from openpyxl import load_workbook
from pathlib import Path

DB = dict(host="localhost", port=5433, dbname="activ_education",
          user="postgres", password="abalakata")

XLSX = Path(__file__).parent.parent / "Base_Etablissements_Superieurs_Togo.xlsx"

def main():
    # 1. Lecture XLSX : secteurs par (ID_ETABLISSEMENT)
    wb = load_workbook(XLSX, data_only=True)
    ws = wb["Debouches"]
    secteurs_par_etab = defaultdict(set)
    for row in ws.iter_rows(min_row=2, values_only=True):
        id_form, id_etab, formation, metiers, secteur, salaire, persp = row[:7]
        if id_etab and secteur and secteur != "Non disponible":
            secteurs_par_etab[id_etab].add(secteur.strip())

    print(f"  {len(secteurs_par_etab)} établissements XLSX avec secteurs")

    # 2. Mapping XLSX ID → fiche_id
    conn = psycopg2.connect(**DB)
    conn.autocommit = False
    cur = conn.cursor()
    cur.execute("SELECT fiche_id, id_source_xlsx FROM etablissement_details_xlsx")
    xlsx_to_fiche = {r[1]: r[0] for r in cur.fetchall()}

    secteurs_par_fiche_id = {}
    for xlsx_id, secteurs in secteurs_par_etab.items():
        fiche_id = xlsx_to_fiche.get(xlsx_id)
        if fiche_id:
            secteurs_par_fiche_id[fiche_id] = secteurs

    print(f"  {len(secteurs_par_fiche_id)} établissements avec fiche_id mappé")

    # 3. Pour chaque filière : secteurs cumulés via ses établissements liés
    cur.execute("""
        SELECT ef.filiere_id, ef.etablissement_id
        FROM etablissement_filiere ef
    """)
    filiere_etabs = defaultdict(set)
    for fil_id, etab_id in cur.fetchall():
        filiere_etabs[fil_id].add(etab_id)

    filiere_secteurs = defaultdict(set)
    for fil_id, etab_ids in filiere_etabs.items():
        for etab_id in etab_ids:
            for s in secteurs_par_fiche_id.get(etab_id, []):
                filiere_secteurs[fil_id].add(s)

    print(f"  {len(filiere_secteurs)} filières avec secteurs connus")

    # 4. Toutes les fiches_metier avec leur secteur
    cur.execute("""
        SELECT fm.id, fm.secteur
        FROM fiches_metier fm
        WHERE fm.secteur IS NOT NULL AND fm.secteur <> ''
    """)
    metiers_par_secteur = defaultdict(list)
    for m_id, secteur in cur.fetchall():
        metiers_par_secteur[secteur.strip().lower()].append(m_id)

    print(f"  {len(metiers_par_secteur)} secteurs distincts côté métier")

    # 5. Matching
    inserts = []
    for fil_id, secteurs in filiere_secteurs.items():
        metiers_match = set()
        for s in secteurs:
            metiers_match.update(metiers_par_secteur.get(s.lower(), []))
        for m_id in metiers_match:
            inserts.append((fil_id, m_id))

    print(f"  {len(inserts)} liens filière↔métier à insérer")

    # 6. INSERT idempotent
    cur.executemany("""
        INSERT INTO filiere_metier (filiere_id, metier_id)
        VALUES (%s, %s)
        ON CONFLICT DO NOTHING
    """, inserts)
    conn.commit()

    cur.execute("SELECT COUNT(*) FROM filiere_metier")
    print(f"  Total liens filière↔métier en base : {cur.fetchone()[0]}")

    cur.close()
    conn.close()

if __name__ == "__main__":
    main()